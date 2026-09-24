import os
from datetime import datetime
from django.core.management.base import BaseCommand
from django.db.models import Q
from django.urls import reverse, NoReverseMatch
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from customer_dashboard.models import Customer
from tally_voucher.models import VoucherStockItem


def format_address(customer):
    parts = []
    if customer.address and customer.address.strip():
        parts.append(customer.address.strip())
    if customer.district and customer.district.strip():
        parts.append(customer.district.strip())
    if customer.state and customer.state.strip():
        parts.append(customer.state.strip())
    addr_str = ", ".join(parts)

    pincode = (customer.pincode or "").strip()
    if pincode.endswith(".0"):
        pincode = pincode[:-2]
    if pincode:
        addr_str = f"{addr_str} - {pincode}" if addr_str else pincode
    return addr_str or "-"


def format_phone(customer):
    phone = (customer.phone or "").strip()
    if phone.endswith(".0"):
        phone = phone[:-2]
    return phone or "-"


def format_dashboard_link(customer, base_url=""):
    try:
        rel_url = reverse("customers:customer_detail", args=[customer.id])
    except NoReverseMatch:
        try:
            rel_url = reverse("customer_detail", args=[customer.id])
        except NoReverseMatch:
            rel_url = f"/customers/customer/{customer.id}/"
    return f"{base_url}{rel_url}" if base_url else rel_url


def build_worksheet(ws, title, headers, rows, link_col_idx=None, center_col_indices=None):
    ws.title = title
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_side = Side(style="thin", color="D1D5DB")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border
    ws.row_dimensions[1].height = 28

    body_font = Font(name="Calibri", size=10)
    link_font = Font(name="Calibri", size=10, color="2563EB", underline="single")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for row_idx, r in enumerate(rows, start=2):
        ws.append(r)
        max_lines = max((str(v).count("\n") + 1 for v in r), default=1)
        ws.row_dimensions[row_idx].height = max(22, max_lines * 16)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = border

            if center_col_indices and col_idx in center_col_indices:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

            if link_col_idx and col_idx == link_col_idx:
                cell.font = link_font
                val = cell.value
                if val and val != "-":
                    cell.hyperlink = val
            else:
                cell.font = body_font

    ws.freeze_panes = "A2"

    for col in ws.columns:
        max_len = 0
        for cell in col:
            lines = str(cell.value or "").split("\n")
            line_len = max((len(l) for l in lines), default=0)
            if line_len > max_len:
                max_len = line_len
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(min(max_len + 4, 60), 14)


class Command(BaseCommand):
    help = (
        "Export customers who purchased Erkodur/Erkoflex, Zendura Flex, or both on Tax Invoices, "
        "including exact product names and purchase dates."
    )

    def add_arguments(self, parser):
        parser.add_argument(
            "--output",
            "-o",
            type=str,
            default=None,
            help="Path/Filename for the exported Excel file (.xlsx). Defaults to 'ortho_product_buyers_YYYYMMDD_HHMMSS.xlsx'",
        )
        parser.add_argument(
            "--base-url",
            type=str,
            default="",
            help="Optional base URL (e.g. http://127.0.0.1:8000 or https://oblu.in) to prefix customer dashboard links",
        )

    def handle(self, *args, **options):
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_file = options.get("output") or f"ortho_product_buyers_{timestamp}.xlsx"
        base_url = (options.get("base_url") or "").rstrip("/")

        self.stdout.write(self.style.NOTICE("Querying Tax Invoices for Erkodur/Erkoflex and Zendura products..."))

        # 1. Base Tax Invoice filter
        tax_invoice_filter = Q(voucher__voucher_type__iexact="Tax Invoice")

        # 2. Product Name Filters
        erkodur_q = (
            Q(item__name__icontains="erkodur")
            | Q(item__name__icontains="erkoflex")
            | Q(item_name_text__icontains="erkodur")
            | Q(item_name_text__icontains="erkoflex")
        )

        zendura_q = (
            Q(item__name__icontains="zendura flex")
            | Q(item__name__icontains="zendura flx")
            | Q(item_name_text__icontains="zendura flex")
            | Q(item_name_text__icontains="zendura flx")
        )

        # 3. Aggregate latest date and distinct product names for Erkodur / Erkoflex
        erkodur_items_qs = (
            VoucherStockItem.objects.filter(tax_invoice_filter & erkodur_q)
            .values("voucher__party_name", "voucher__date", "item__name", "item_name_text")
        )

        erkodur_map = {}
        for entry in erkodur_items_qs:
            party = entry.get("voucher__party_name")
            if not party:
                continue
            key = party.strip().lower()
            dt = entry.get("voucher__date")
            pname = (entry.get("item__name") or entry.get("item_name_text") or "").strip()

            if key not in erkodur_map:
                erkodur_map[key] = {"last_date": dt, "products": set()}
            if dt and (erkodur_map[key]["last_date"] is None or dt > erkodur_map[key]["last_date"]):
                erkodur_map[key]["last_date"] = dt
            if pname:
                erkodur_map[key]["products"].add(pname)

        # 4. Aggregate latest date and distinct product names for Zendura Flex
        zendura_items_qs = (
            VoucherStockItem.objects.filter(tax_invoice_filter & zendura_q)
            .values("voucher__party_name", "voucher__date", "item__name", "item_name_text")
        )

        zendura_map = {}
        for entry in zendura_items_qs:
            party = entry.get("voucher__party_name")
            if not party:
                continue
            key = party.strip().lower()
            dt = entry.get("voucher__date")
            pname = (entry.get("item__name") or entry.get("item_name_text") or "").strip()

            if key not in zendura_map:
                zendura_map[key] = {"last_date": dt, "products": set()}
            if dt and (zendura_map[key]["last_date"] is None or dt > zendura_map[key]["last_date"]):
                zendura_map[key]["last_date"] = dt
            if pname:
                zendura_map[key]["products"].add(pname)

        # 5. Segment Parties
        all_erkodur_parties = set(erkodur_map.keys())
        all_zendura_parties = set(zendura_map.keys())
        both_parties = all_erkodur_parties.intersection(all_zendura_parties)

        self.stdout.write(
            f"Tax Invoice stats: {len(all_erkodur_parties)} Erkodur/Erkoflex buyers, "
            f"{len(all_zendura_parties)} Zendura buyers, "
            f"{len(both_parties)} purchased both."
        )

        # 6. Fetch Customers and build data structures
        customers = Customer.objects.select_related("salesperson").all()

        dual_rows = []
        erkodur_rows = []
        zendura_rows = []
        all_rows = []

        for customer in customers:
            if not customer.name:
                continue
            key = customer.name.strip().lower()

            has_erkodur = key in all_erkodur_parties
            has_zendura = key in all_zendura_parties

            if not (has_erkodur or has_zendura):
                continue

            addr_str = format_address(customer)
            phone_str = format_phone(customer)
            salesperson_str = (
                customer.salesperson.name
                if customer.salesperson and customer.salesperson.name
                else "Unassigned"
            )
            dash_link = format_dashboard_link(customer, base_url)

            # Erkodur info
            e_info = erkodur_map.get(key, {})
            e_last_dt = e_info.get("last_date")
            e_products = sorted(e_info.get("products", set()))
            e_prods_str = "\n".join(e_products) if e_products else "-"
            e_date_str = e_last_dt.strftime("%Y-%m-%d") if e_last_dt else "-"

            # Zendura info
            z_info = zendura_map.get(key, {})
            z_last_dt = z_info.get("last_date")
            z_products = sorted(z_info.get("products", set()))
            z_prods_str = "\n".join(z_products) if z_products else "-"
            z_date_str = z_last_dt.strftime("%Y-%m-%d") if z_last_dt else "-"

            # Overall latest date
            dates = [d for d in (e_last_dt, z_last_dt) if d]
            overall_dt = max(dates) if dates else None
            overall_date_str = overall_dt.strftime("%Y-%m-%d") if overall_dt else "-"

            # Combined product list
            all_prods = sorted(set(e_products).union(set(z_products)))
            all_prods_str = "\n".join(all_prods) if all_prods else "-"

            # 6a. Tab 1: Dual Buyers (Both)
            if has_erkodur and has_zendura:
                dual_rows.append({
                    "row": [
                        customer.name,
                        addr_str,
                        phone_str,
                        salesperson_str,
                        e_prods_str,
                        e_date_str,
                        z_prods_str,
                        z_date_str,
                        overall_date_str,
                        dash_link,
                    ],
                    "_sort": overall_dt,
                })

            # 6b. Tab 2: Erkodur & Erkoflex Buyers
            if has_erkodur:
                erkodur_rows.append({
                    "row": [
                        customer.name,
                        addr_str,
                        phone_str,
                        salesperson_str,
                        e_prods_str,
                        e_date_str,
                        "Yes" if has_zendura else "No",
                        dash_link,
                    ],
                    "_sort": e_last_dt,
                })

            # 6c. Tab 3: Zendura Buyers
            if has_zendura:
                zendura_rows.append({
                    "row": [
                        customer.name,
                        addr_str,
                        phone_str,
                        salesperson_str,
                        z_prods_str,
                        z_date_str,
                        "Yes" if has_erkodur else "No",
                        dash_link,
                    ],
                    "_sort": z_last_dt,
                })

            # 6d. Tab 4: All Target Customers
            category_tag = (
                "Both (Dual Buyer)"
                if (has_erkodur and has_zendura)
                else ("Erkodur / Erkoflex Only" if has_erkodur else "Zendura Only")
            )
            all_rows.append({
                "row": [
                    customer.name,
                    category_tag,
                    addr_str,
                    phone_str,
                    salesperson_str,
                    all_prods_str,
                    overall_date_str,
                    dash_link,
                ],
                "_sort": overall_dt,
            })

        # Sort descending by latest purchase date
        dual_rows.sort(key=lambda r: r["_sort"], reverse=True)
        erkodur_rows.sort(key=lambda r: r["_sort"], reverse=True)
        zendura_rows.sort(key=lambda r: r["_sort"], reverse=True)
        all_rows.sort(key=lambda r: r["_sort"], reverse=True)

        self.stdout.write(
            f"Matched Customers: {len(dual_rows)} Dual Buyers, "
            f"{len(erkodur_rows)} Erkodur/Erkoflex buyers, "
            f"{len(zendura_rows)} Zendura buyers, "
            f"{len(all_rows)} Total unique customers."
        )

        # 7. Create Excel Workbook with 4 organized tabs
        wb = openpyxl.Workbook()

        # --- Sheet 1: Dual Buyers (Both) ---
        ws_dual = wb.active
        headers_dual = [
            "Customer Name",
            "Address",
            "Phone Number",
            "Salesperson Handle / Name",
            "Erkodur / Erkoflex Products Bought",
            "Last Bought (Erkodur / Erkoflex) Date",
            "Zendura Products Bought",
            "Last Bought (Zendura Flex) Date",
            "Overall Latest Date across these products",
            "Customer Dashboard Link",
        ]
        build_worksheet(
            ws_dual,
            "Dual Buyers (Both)",
            headers_dual,
            [r["row"] for r in dual_rows],
            link_col_idx=10,
            center_col_indices=(3, 6, 8, 9),
        )

        # --- Sheet 2: Erkodur & Erkoflex Buyers ---
        ws_erkodur = wb.create_sheet()
        headers_erkodur = [
            "Customer Name",
            "Address",
            "Phone Number",
            "Salesperson Handle / Name",
            "Erkodur / Erkoflex Products Bought",
            "Last Bought Date",
            "Also Bought Zendura?",
            "Customer Dashboard Link",
        ]
        build_worksheet(
            ws_erkodur,
            "Erkodur & Erkoflex Buyers",
            headers_erkodur,
            [r["row"] for r in erkodur_rows],
            link_col_idx=8,
            center_col_indices=(3, 6, 7),
        )

        # --- Sheet 3: Zendura Buyers ---
        ws_zendura = wb.create_sheet()
        headers_zendura = [
            "Customer Name",
            "Address",
            "Phone Number",
            "Salesperson Handle / Name",
            "Zendura Products Bought",
            "Last Bought Date",
            "Also Bought Erkodur/Erkoflex?",
            "Customer Dashboard Link",
        ]
        build_worksheet(
            ws_zendura,
            "Zendura Buyers",
            headers_zendura,
            [r["row"] for r in zendura_rows],
            link_col_idx=8,
            center_col_indices=(3, 6, 7),
        )

        # --- Sheet 4: All Target Customers (Consolidated) ---
        ws_all = wb.create_sheet()
        headers_all = [
            "Customer Name",
            "Purchase Category",
            "Address",
            "Phone Number",
            "Salesperson Handle / Name",
            "All Matching Products Bought",
            "Overall Latest Date",
            "Customer Dashboard Link",
        ]
        build_worksheet(
            ws_all,
            "All Target Customers",
            headers_all,
            [r["row"] for r in all_rows],
            link_col_idx=8,
            center_col_indices=(2, 4, 7),
        )

        wb.save(output_file)
        self.stdout.write(
            self.style.SUCCESS(
                f"\nSuccessfully generated multi-tab Excel report: {output_file}\n"
                f" - Tab 1: 'Dual Buyers (Both)' ({len(dual_rows)} customers)\n"
                f" - Tab 2: 'Erkodur & Erkoflex Buyers' ({len(erkodur_rows)} customers)\n"
                f" - Tab 3: 'Zendura Buyers' ({len(zendura_rows)} customers)\n"
                f" - Tab 4: 'All Target Customers' ({len(all_rows)} customers)\n"
            )
        )
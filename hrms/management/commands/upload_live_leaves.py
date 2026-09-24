import os
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
import openpyxl

from hrms import models as m


class Command(BaseCommand):
    help = "Loads leave balances directly into EmployeeLeaveBalance from an Excel file."

    def add_arguments(self, parser):
        parser.add_argument('excel_file', type=str, help='Path to the Excel file (.xlsx)')

    def _parse_val(self, val, default=0.0):
        if val is None:
            return default
        cleaned = str(val).strip()
        if cleaned in ('', '-', '--', 'N/A', 'none', 'null'):
            return default
        try:
            return float(cleaned)
        except (ValueError, TypeError):
            return default

    def handle(self, *args, **options):
        file_path = options['excel_file']

        if not os.path.exists(file_path):
            raise CommandError(f"File not found: {file_path}")

        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active

        # Build header index mapping
        headers = {}
        for col_idx in range(1, sheet.max_column + 1):
            val = sheet.cell(row=1, column=col_idx).value
            if val:
                headers[str(val).strip().lower().replace('_', ' ')] = col_idx

        def find_col(candidates):
            for c in candidates:
                if c in headers:
                    return headers[c]
            return None

        id_col = find_col(['employee id', 'employee code', 'emp id', 'empid', 'code'])
        if not id_col:
            raise CommandError(f"Could not find 'Employee Id' column. Found headers: {list(headers.keys())}")

        col_bereave = find_col(['bereave', 'bereavement', 'bl'])
        col_menstrual = find_col(['menstrual', 'ml'])
        col_sick = find_col(['sick', 'sl'])
        col_earned = find_col(['earned', 'el'])
        col_casual = find_col(['casual', 'cl'])
        col_compoff = find_col(['comp off', 'comp-off', 'compoff', 'co'])

        updated = 0
        skipped = []

        with transaction.atomic():
            for row_idx in range(2, sheet.max_row + 1):
                raw_code = sheet.cell(row=row_idx, column=id_col).value
                if not raw_code or not str(raw_code).strip():
                    continue

                emp_code = str(raw_code).strip()
                emp = m.Employee.objects.filter(employee_code__iexact=emp_code).first()

                if not emp:
                    skipped.append(emp_code)
                    continue

                # Parse row values
                bereave = self._parse_val(sheet.cell(row=row_idx, column=col_bereave).value) if col_bereave else 0.0
                menstrual = self._parse_val(sheet.cell(row=row_idx, column=col_menstrual).value) if col_menstrual else 0.0
                sick = self._parse_val(sheet.cell(row=row_idx, column=col_sick).value) if col_sick else 0.0
                earned = self._parse_val(sheet.cell(row=row_idx, column=col_earned).value) if col_earned else 0.0
                casual = self._parse_val(sheet.cell(row=row_idx, column=col_casual).value) if col_casual else 0.0
                compoff = self._parse_val(sheet.cell(row=row_idx, column=col_compoff).value) if col_compoff else 0.0

                # Update EmployeeLeaveBalance
                bank, _ = m.EmployeeLeaveBalance.objects.get_or_create(e_name=emp)
                bank.bereavement_leave = bereave
                bank.menstrual_leave = menstrual
                bank.sick_leave = sick
                bank.earned_leave = earned
                bank.casual_leave = casual
                bank.comp_off = compoff
                bank.save()

                updated += 1

        self.stdout.write(self.style.SUCCESS(f"Successfully fed balances for {updated} employee(s) into EmployeeLeaveBalance."))
        if skipped:
            self.stdout.write(self.style.WARNING(f"Skipped {len(skipped)} unmatched code(s): {', '.join(skipped)}"))